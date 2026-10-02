import configparser
import os
import requests
from datetime import datetime
import openpyxl

def build_export_url(url, sheet_name=None):
    if not url or str(url).strip() == "":
        return None

    if '/d/' in url:
        base_id = url.split('/d/')[1].split('/')[0]
        export_url = f"https://docs.google.com/spreadsheets/d/{base_id}/export?format=xlsx"
        
        if sheet_name and str(sheet_name).strip() != "":
            sn_clean = str(sheet_name).strip()
            if sn_clean.isdigit():
                export_url += f"&gid={sn_clean}"
            else:
                export_url += f"&sheet={sn_clean}"
                
        return export_url
    return None

def sanitize_excel_headers(filename):
    try:
        wb = openpyxl.load_workbook(filename)
        for sheet in wb.worksheets:
            seen_headers = {}
            for col_idx in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(row=1, column=col_idx).value
                if cell_value is not None:
                    val_str = str(cell_value).strip()
                    val_upper = val_str.upper()
                    
                    if val_upper in seen_headers:
                        seen_headers[val_upper] += 1
                        sheet.cell(row=1, column=col_idx).value = f"{val_str} {seen_headers[val_upper]}"
                    else:
                        seen_headers[val_upper] = 1

            table_names = list(sheet.tables.keys())
            for tbl_name in table_names:
                del sheet.tables[tbl_name]

        wb.save(filename)
        wb.close()
        print(f"--> Sanitasi header & tabel pada {filename} berhasil dilakukan.")
    except Exception as e:
        print(f"--> Peringatan: Tidak dapat melakukan sanitasi pada {filename}: {e}")

def download_file(url, filename):
    try:
        print(f"--> Sedang mengunduh file: {filename}...")
        response = requests.get(url)
        response.raise_for_status()
        
        with open(filename, 'wb') as file:
            file.write(response.content)
            
        print(f"--> File berhasil disimpan: {filename}")
        
        sanitize_excel_headers(filename)
        
    except Exception as e:
        print(f"--> Terjadi kesalahan saat mengunduh {filename}: {e}")

def main():
    config_file = 'config.conf'
    
    if not os.path.exists(config_file):
        print(f"--> Peringatan: File '{config_file}' tidak ditemukan!")
        return

    config = configparser.ConfigParser()
    config.read(config_file)

    if not config.has_section('SS'):
        print("--> Section [SS] tidak ditemukan di config.conf!")
        return

    current_date = datetime.now().strftime("%d-%m-%Y")

    targets = [
        {
            "url_key": "url-irc",
            "sn_key": "url-irc-sn",
            "filename": f"ORDER IRC JATENG {current_date}_temp.xlsx"
        },
        {
            "url_key": "url-zn",
            "sn_key": "url-zn-sn",
            "filename": f"ORDER ZN JATENG {current_date}_temp.xlsx"
        }
    ]

    for target in targets:
        url_key = target["url_key"]
        sn_key = target["sn_key"]
        filename = target["filename"]

        raw_url = config.get('SS', url_key, fallback=None)
        sheet_name = config.get('SS', sn_key, fallback=None)

        if not raw_url:
            print(f"--> Peringatan: '{url_key}' kosong atau tidak ditemukan di config.conf. Dilewati.")
            continue

        download_url = build_export_url(raw_url, sheet_name)
        
        if download_url:
            download_file(download_url, filename)
        else:
            print(f"--> URL tidak valid untuk '{url_key}': {raw_url}")

if __name__ == "__main__":
    print("--> Memulai proses pengunduhan...")
    main()
    print("--> Semua proses selesai.")